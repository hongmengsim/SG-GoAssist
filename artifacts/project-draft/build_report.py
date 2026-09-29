from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output/project-draft'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.66)
sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.3)
styles = doc.styles
for border in styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(7)
styles['Normal'].font.size=Pt(10)
styles['Normal'].paragraph_format.line_spacing=1.5
styles['Title'].font.size=Pt(10)
styles['Title'].font.bold=True
styles['Heading 1'].font.size=Pt(10)
styles['Heading 1'].font.bold=True
styles['Heading 1'].paragraph_format.space_after=Pt(12)
styles['Heading 2'].font.size=Pt(10)
styles['Heading 2'].font.bold=True
styles['Heading 2'].paragraph_format.space_before=Pt(12)
styles['Heading 2'].paragraph_format.space_after=Pt(5)
styles['Caption'].font.size=Pt(10)
styles['Subtitle'].font.size=Pt(10)
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].paragraph_format.line_spacing=1.5
footer = sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('SG GoAssist  |  Draft  |  '); r.font.name='Arial'; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string('666666')
footer.paragraph_format.line_spacing=1.5
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
doc.core_properties.title='SG GoAssist Project Report Draft'
doc.core_properties.subject='Autonomous Passenger Assistance System project and progress'
doc.core_properties.author='SmaRHt Buses'

def p(text='',style=None):
    text=re.sub(r'\[[0-9, –-]+\]',lambda m:m.group(0).replace(' ','\u00a0'),text)
    return doc.add_paragraph(text,style)
def h(text):
    heading=doc.add_heading(text,2)
    if text=='Conclusion': heading.paragraph_format.page_break_before=True
    return heading
def page(title):
    if title in ['Abstract','1 Introduction','11 References']:
        doc.add_page_break()
    heading=doc.add_heading(title,1)
    heading.paragraph_format.space_before=Pt(16)
def note(text):
    a=p(text); a.paragraph_format.space_after=Pt(8)
    for r in a.runs: r.font.size=Pt(10); r.font.italic=True; r.font.color.rgb=RGBColor.from_string('666666')
def blank(label,lines=2):
    h(label); note('To complete')
    for _ in range(lines):
        a=p(); a.paragraph_format.space_after=Pt(10)
def table(headers,rows,widths):
    widths=[w*6.27/sum(widths) for w in widths]
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,x in enumerate(headers): t.rows[0].cells[i].text=x
    for data in rows:
        cells=t.add_row().cells
        for i,x in enumerate(data): cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'243F50' if ri==0 else ('F2F5F7' if ri%2 else 'FFFFFF')); tcpr.append(shade)
            margin=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'65'); e.set(qn('w:type'),'dxa'); margin.append(e)
            tcpr.append(margin)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1); para.paragraph_format.space_before=Pt(1); para.paragraph_format.line_spacing=1.5
                para.paragraph_format.keep_with_next=(ri==0 or (headers[0]=='Model record' and ri<len(t.rows)-1))
                for run in para.runs:
                    run.font.name='Arial'; run.font.size=Pt(10); run.font.bold=(ri==0)
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    borders=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+e); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    p().paragraph_format.space_after=Pt(0)
    return t

p('SG GoAssist Autonomous Passenger Assistance System', 'Title')
p('Draft project report', 'Subtitle')
p('Singapore BusTech Grand Challenge 2026\nIHL Student Category\nProgress as at 25 September 2026')
p('Institution  National University of Singapore')
p('Team name  SmaRHt Buses')
p('Team members\nJin Donghua\nLiu Xuan Tong\nRyan Koh Jun Hao\nSim Hong Meng\nZeph Tang Zhenkai')
p('Staff in charge  Koenraad Mouthaan')
p('Report submission deadline  30 September 2026')
note('Draft for group completion. Incomplete results and budget fields remain blank.')
page('Abstract')
p('SG GoAssist aims to help passengers request and receive boarding or alighting assistance on an autonomous bus. Our prototype combines an accessible request interface, local perception, a deterministic assistance controller and a model ramp. The intended outcome is a complete journey in which the system acknowledges a request, checks safety conditions, provides assistance and verifies that the ramp is stowed before departure. [1, 2]')
p('The current project has software and firmware for this assistance loop, a detailed bus and ramp CAD model, and an ESP32-S3 bench setup that has produced valid distance readings. The software includes passenger and operator interfaces, assistance-case management, telemetry, interlocks and audit logging. CAD records document a discrete motion-clearance study. These are separate forms of progress and do not yet establish a completed physical system. [3–7]')
p('The next priority is physical integration and controlled validation. Ramp actuation, obstacle coverage, model accuracy, repeatability and user experience still require measured results. The results and feedback sections below remain open until those activities are complete. The report follows the four-week prototype plan and its D1 to D9 deliverables. [1]')
h('Report contents')
p('1 Introduction\n2 Design alternatives and original contributions\n3 Methods and system design\n4 Software implementation results\n5 Mechanical and electronics results\n6 Machine learning methods and status\n7 Deliverable readiness\n8 Validation methods and pending results\n9 Discussion and remaining work\n10 Proposed budget and team records\n11 References')

page('1 Introduction')
h('Problem and intended users')
p('The project addresses the assistance tasks that remain around an autonomous bus journey: identifying the correct service, requesting enough boarding time, accessing a ramp and understanding whether help is on its way. A passenger standing near a bus stop may be waiting for another service, so presence alone cannot establish boarding intent. Visible mobility aids also cannot represent every assistance need. [1, 2]')
p('The proposed system supports wheelchair and personal mobility aid users, passengers using walking aids, older passengers, people with temporary or non-visible needs, and passengers who benefit from audio, visual or haptic information. These are intended user groups. Representative-user validation remains an unfinished project activity. [1, 2]')
h('Project objectives')
table(['Objective','Planned evidence'],[
('Recognise observable assistance cues','Held-out class metrics and examples of errors for the perception models.'),
('Act on explicit passenger intent','Correct request acknowledgement, bus assignment and cancellation in controlled scenarios.'),
('Provide a verified mechanical response','Safe model-ramp deployment, boarding hold and confirmed retraction.'),
('Measure accessibility outcomes','Comparison with a manual-assistance baseline and consented participant feedback.')],[2.15,4.6])
h('Minimum viable prototype')
p('The action plan defines one repeatable journey using a stop unit or lightweight app, a mock bus, a boarding fixture, safety sensors and an operator dashboard. A second scenario requests additional time and audio or visual guidance without a ramp. This demonstrates that an explicit request can support needs that a camera cannot infer. [1]')
p('The prototype scope excludes a road-capable autonomous vehicle, a full-scale certified ramp and claims of reliable perception in every public environment. Production transport feeds, city-wide infrastructure and advanced walking guidance are outside the core demonstration. The repository contains additional journey-planning and assistant features, but these do not replace the core integration and validation tasks. [1, 4]')
h('Development of the concept')
p('The initial expression of interest proposed RFID requests, camera/LiDAR recognition and pressure or magnetic sensing. The later design gives a tactile stop button a central role, while the action plan retains an app and optional NFC. The implemented software accepts several input types under one assistance-case contract. For the first physical demonstration, the team still needs to freeze the actual input devices and sensor configuration. [1–3, 9]')

page('2 Design alternatives and original contributions')
h('Alternatives and selection rationale')
p('The design progressed from the expression of interest to the later request-driven prototype. The alternatives below explain the current direction without treating earlier concepts as completed hardware. [1, 2, 9]')
table(['Alternative','Tradeoff','Selected approach'],[
('RFID or NFC only','Simple explicit request, but depends on obtaining and carrying a card.','Optional input alongside a tactile button and app.'),
('Camera-triggered ramp','Recognises observable cues, but cannot establish which bus a passenger intends to board.','Use perception for cues and dwell; require confirmed ramp intent.'),
('Pressure or magnetic presence sensing','Adds redundancy but can confuse passengers, luggage and passers-by.','Treat presence as supporting evidence, with geometric safety checks.'),
('App-only request','Flexible selection and feedback, but excludes passengers without a suitable device.','Keep a stop-based tactile request route.'),
('Full-scale vehicle actuation','Closer to deployment, with major mechanical and safety demands.','Demonstrate on a low-voltage mock bus with measured interlocks.')],[1.5,2.55,2.7])
h('Team contribution and use of existing technology')
p('The team contribution is the integration of explicit passenger intent with an assistance-case workflow, deterministic safety checks, local actuator verification and accessible feedback. The repository implements the request and verification logic, and the CAD work develops the centre-door boarding mechanism concept. These are the elements to demonstrate and explain as project work. [2, 3, 6]')
p('ESP32 hardware, ToF sensors, RFID/NFC, ArUco pose estimation and pretrained model architectures are existing technologies. The bus model is reference-inspired, and proposed mechanisms are not OEM-verified. The team should credit these foundations and identify any adapted code, models or designs in the final references. No claim of worldwide novelty or patentability is made. [6, 8, 9]')
p('The official rubric allocates 20% to design and methodology, 30% to execution, 30% to originality and creativity, and 20% to project presentation. Only original elements of an improvement on an existing design receive originality credit. Final contribution records and comparisons should therefore identify precisely what the group designed and implemented. [11, p. 16]')

page('3 Methods and system design')
h('System architecture')
table(['Layer','Role in the prototype'],[
('Passenger and stop unit','App, tactile button or optional NFC records the requested assistance. Status returns through accessible feedback.'),
('Local perception','An edge computer interprets camera observations. The VL53L0X reports one distance within its monitored field.'),
('Assistance controller','The backend retains intent, assigns the bus and stop, checks capability and manages the assistance case.'),
('Mock bus controller','ESP32 firmware receives commands, repeats local safety checks and reports actuator and limit-switch status.'),
('Operator and evidence','The console supports confirmation and fault recovery. Metrics and audit events record the outcome.')],[1.8,4.95])
h('Intended demonstration sequence')
p('A passenger requests ramp assistance and, where needed, extra time. The system acknowledges the request for the selected bus. Perception can support positioning or dwell decisions, while the request remains authoritative. The model bus then docks at the marked boarding fixture and confirms that it is stopped securely. [1, 3]')
p('The ramp area is visually marked. Before deployment, the controller checks the bus and stop match, available capability and space, fresh telemetry, stopped and brake signals, open door, clear path and controller health. The local controller repeats relevant checks while the ramp moves. “Ready” requires both completed actuation and a deployed limit-switch indication. [3, 5]')
p('The system maintains assistance during boarding. A completion or cancellation event initiates a safe closure sequence. A case reaches its terminal state only after stowed-ramp verification. Departure also requires a stowed ramp and clear departure conditions. Faults or ambiguous states remain visible to the operator. [3, 5]')
h('Safety and passenger control')
p('Sensor-only detections cannot confirm a ramp request. They enter a confirmation path, while reversible extra dwell may proceed from a detection. Missing or stale safety data blocks movement. The prototype design also requires a physical emergency stop and independent electrical interlocks before physical demonstrations. Their installation and behavior must be checked on the assembled model. [2, 3, 5]')
h('Privacy boundary')
p('The reference edge observer processes frames locally and sends class/confidence metadata. It does not save or transmit camera frames. The system is intended to recognise objects and assistance cues without identifying people or diagnosing disabilities. Any future remote-video function from the earlier concept requires a separate privacy and consent design. [2, 3, 8]')

page('4 Software implementation results')
h('Implemented software scope')
p('The repository provides a passenger application, a backend and shared message contracts. The assistance controller connects explicit requests, observations, vehicle capabilities, safety telemetry and actuator status. It preserves separate passenger intents while using idempotent commands to avoid duplicate actuator actions. [3, 4]')
table(['Area','Current implementation','Remaining evidence'],[
('Passenger app','Requests, status, cancellation, nearby stops, journeys and accessibility options.','Physical-device and representative-user checks.'),
('Operator console','Case queue, safety state, confirmation, escalation and retry.','Operator recovery trials with the connected model.'),
('Assistance backend','Case states, safety gates, command status, persistence and audit events.','Timestamped integrated hardware runs.'),
('Mock autonomy','Route-state simulation and marker/ToF docking integration points.','Measured physical docking performance.'),
('Automated checks','Repository tests cover request handling, interlocks, intent, persistence and integration paths.','Attach a dated release-test run and its build revision.')],[1.3,2.9,2.55])
h('Evidence boundaries')
p('The readiness matrix records automated software checks and separately labels competition fixtures. Demonstration arrivals, selected stop amenities, shelter information and advisories must retain their fixture labels. They are not evidence of a live transport-provider integration. [4]')
p('The multilingual assistant includes deterministic handling and an optional on-device Qwen path. Connected ARM64-device verification and physical speech-input testing remain release gates. The assistant can interpret requests and explain information, but safety controllers retain authority over ramp, door and vehicle actions. [4]')
h('Software baseline for the demonstration')
p('The final release should identify a fixed source revision and configuration, reproduce the nominal and fault scenarios, and archive the console/event output. Software checks support implementation confidence, but physical trials must establish the behavior of the assembled prototype.')
blank('Release revision and dated test summary',1)

page('5 Mechanical and electronics results')
h('Mechanical model')
p('Robobus v2 includes a centre-door access concept, plug-and-slide door leaves, a translating and tilting ramp, a bridge and a kerb fixture. Native Fusion and STEP exports exist for travel and deployed states. The CAD validation record reports 189 valid solid bodies and eight native as-built joints. [6]')
doc.add_picture(str(ROOT/'cad/Punggol_WeRide_Robobus_v2_deployed.png'),width=Inches(5.1))
p('Figure 1  CAD boarding-ready concept. Proposed geometry, not a photograph of the assembled prototype. [6]', 'Caption')
p('The recorded clearance study sampled 73 poses with zero moving-part intersection volumes reported. A separate passage check addressed a static doorway obstruction. This is a discrete CAD packaging result. It does not verify continuous collision clearance, structural strength, load capacity or physical reliability. [6]')
h('Electronics bench progress')
p('On 23 September, the ESP32-S3 test sketch compiled and uploaded. The connected sensor was identified as a VL53L0X. After initial invalid readings and a full power cycle, the recorded run produced 27 consecutive valid samples from 7639 to 15095 ms, spanning 30 to 136 mm. This establishes live ranging only. No known-distance accuracy result or automatic ramp-stop result follows from that test. [7]')
p('The selected ToF sensor is the VL53L0X. It uses an invisible 940 nm laser and reports one distance over I²C, with a manufacturer-specified range up to 2 m under suitable conditions. This limit is not a measured prototype result. It provides no depth map or object classification, and a single clear reading cannot establish that the entire ramp area is clear. The separate RYS1230 lasers visibly mark the deployment area. [7, 10, 12]')
p('The proposed obstacle check compares valid, fresh VL53L0X readings with a calibrated threshold in its monitored field. Invalid or stale readings must block movement. Controller adaptation, mounting and threshold validation remain incomplete. The bench sketch only reports ranges and keeps the visible laser outputs off. [7]')

page('6 Machine learning methods and status')
h('Planned model functions')
table(['Module','Bounded decision','Progress and next evidence'],[
('ML 1 assistance cues','Detect observable mobility aids and relevant waiting-zone cues.','Training/evaluation scripts and label guidance exist. Complete dataset, exported weights and held-out results.'),
('ML 2 ramp zone','Combine camera classifications with VL53L0X distance checks in its limited monitored field.','Fusion interfaces exist. Adapt and validate single-zone threshold logic; wider ramp coverage requires separate evidence.'),
('ML 3 boarding state','Estimate active, complete or uncertain boarding from temporal features.','The action plan defines the model and a fixed-dwell fallback. Complete sequence labels, training and evaluation.'),
('Docking','Estimate pose relative to a marker and verify range.','ArUco/ToF observer code exists. Calibrate and measure repeated physical docking.')],[1.32,2.25,3.18])
p('The planned models support limited decisions. A cue detector cannot diagnose a disability or authorise ramp movement. ML 2 must operate within the deterministic safety boundary, and ML 3 must retain a sensor veto and a fallback dwell rule. ArUco docking is a marker-based vision method rather than a trained ML model. [1, 5, 8]')
h('Dataset and evaluation method')
p('Use consented or appropriately licensed imagery, document class definitions and split recordings by session or location. Keep the test split separate from training and threshold selection. Record the dataset version, model checksum, test conditions and device used for timing. The ML 1 package includes tools to record provenance and compute precision, recall, mAP and inference speed. [1, 8]')
p('The action plan proposed 300–500 ML 1 images, while the later repository guide proposes 800–1,200. These are planning targets, not dataset counts. The final dataset scope and achieved coverage should be documented below. [1, 8]')
table(['Model record','ML 1','ML 2','ML 3'],[
('Dataset version and split','','',''),('Selected model and weights','','',''),('Test sample count','','',''),('Precision and recall or state accuracy','','',''),('Latency or inference speed','','',''),('Known failure cases','','','')],[2.5,1.42,1.42,1.41])
note('Blank cells are reserved for measured model results and completed model records.')

page('7 Deliverable readiness')
p('The following assessment follows D1 to D9 in the prototype action plan. “Implemented” describes an available software/design artifact. It does not imply physical integration or completion of acceptance trials. [1, 3–8]')
table(['ID and deliverable','Current position','Completion requirement'],[
('D1 Architecture and interfaces','Documented and implemented in shared contracts and controller logic.','Freeze hardware-specific interfaces and the demonstration configuration.'),
('D2 Integrated physical prototype','Firmware and CAD exist; ToF bench ranging verified.','Assemble bus, stop unit and fixture; record 10 cold-start journeys.'),
('D3 Passenger app and dashboard','Implemented, with automated checks recorded in project documentation.','Capture the connected request, status, cancellation and escalation journey.'),
('D4 ML 1 package','Training/evaluation pipeline and labels prepared.','Dataset manifest, trained weights, metrics, latency and model card.'),
('D5 ML 2 and safety fusion','Fusion interfaces exist; VL53L0X controller adaptation remains pending.','Camera results plus calibrated distance, obstruction and fault trials.'),
('D6 ML 3 or fallback','Model and fallback defined in the plan.','Validate the selected dwell method and record premature-clear outcomes.'),
('D7 Mechanical and electrical','CAD/export package, wiring drawings and sensor bring-up available.','As-built wiring, BOM, power budget, actuator/interlock and cycle tests.'),
('D8 Validation and accessibility','Scenario matrix and metric collection paths defined.','Measured baseline comparison, limitations and consented feedback.'),
('D9 Submission package','Draft report and presentation prepared.','Final evidence, compliant media, source/model archive and rehearsal.')],[1.58,2.57,2.6])
h('Scope decisions to close')
p('Confirm VL53L0X mounting, the actuator arrangement, projected-line placement and whether automatic driving remains in the demonstration. The action plan permits scope reductions, including NFC and automatic driving, while preserving the request path, ramp-zone assessment, interlocks, physical ramp and complete evidence trail. [1]')
p('The earlier proposal and later firmware also differ in their treatment of light debris and motion resumption. Until the selected rule is documented and physically validated, an occupied or uncertain ramp zone should remain a blocking demonstration condition. [1, 2, 5]')

page('8 Validation methods and pending results')
h('Controlled test protocol')
p('Freeze the build, wiring and calibration before testing. Compare VL53L0X readings with known distances and test obstacles inside and outside its monitored field to expose blind spots. Test nominal boarding and alighting, cancellation, duplicate requests, ambiguous perception, sensor failure, network loss and non-ramp assistance. Record expected and actual responses, timestamps and interventions, and retain audit logs. A clear range alone cannot pass a whole-ramp clearance test. [1, 3, 5, 7]')
p('For the baseline, repeat the same scripted journey with a person handling the assistance request, then with SG GoAssist. Compare acknowledgement time, successful completion, manual interventions, unnecessary waiting and perceived confidence. Participant counts, task conditions and failures must accompany the comparison. [1]')
h('Acceptance targets and measured outcomes')
table(['Measure','Proposed acceptance target','Measured result'],[
('Request delivery','20/20 local trials in the action plan; later traceability calls for 30/30 per explicit source.',''),
('Response time','Within 2 s in the action plan; later traceability specifies P95 < 2 s.',''),
('ML 1','At least 90% controlled wheelchair/PMA recall and at least 10 FPS.',''),
('Obstacle and interlock tests','Zero missed defined obstructions; zero unsafe movements in at least 50 interlock/obstruction trials.',''),
('Actuation','50 ramp/door cycles without binding or interlock bypass.',''),
('Docking','10 consecutive trials within demonstrated ramp reach.',''),
('Boarding/dwell','No premature CLEAR in defined sequence tests; fallback available.',''),
('Complete journey','At least 9 of 10 nominal cold-start journeys without manual intervention.','')],[1.28,3.92,1.55])
note('Targets are drawn from the action plan and later traceability document. Blank results are intentional. Reconcile the request-trial count and latency definition in the final test protocol. [1, 5]')
blank('Measured baseline comparison',1)
blank('Integrated demonstration results and failure analysis',1)

page('9 Discussion and remaining work')
h('Analysis of current results')
p('The CAD study supports packaging feasibility at sampled positions, and the VL53L0X log shows communication and ranging after recovery. Neither result demonstrates a safe complete journey. Limited sensing coverage, pending controller adaptation, untested actuation and incomplete model evidence prevent a reliable estimate of end-to-end performance. Measured user benefit remains an open research outcome. [6, 7]')
h('Immediate integration work')
p('First, adapt the controller to VL53L0X single-zone readings and calibrate its distance threshold and monitored field. Validate invalid-reading and stale-data blocking. Connect and verify the ramp actuator, driver, power supply, position feedback, request input and independent stop/interlock functions. Then connect physical telemetry and actuator results to the assistance case. [1, 5, 7]')
p('Next, reproduce one nominal journey before expanding the scenario matrix. Validate the selected perception models and dwell method on held-out data and the actual demonstration device. Use the fixed extended-dwell fallback if ML 3 cannot meet its evidence requirements. The final stage is repeatability testing, evidence capture and rehearsal. [1]')
table(['Risk','Effect on the demonstration','Planned response'],[
('Single-zone blind spots','A valid VL53L0X range does not cover the whole ramp area.','Map the monitored field and test blind spots; retain independent interlocks.'),
('Actuator or mechanism fault','Ramp fails to reach or leave its intended position.','Verify limits, power, mechanical clearance and safe operator recovery.'),
('Insufficient ML evidence','Recognition claims cannot be substantiated.','Reduce classes if needed and report held-out results and failure cases.'),
('Late integration','Subsystem progress does not produce a complete journey.','Freeze interfaces, integrate the core path and defer peripheral features.'),
('Overstated impact','Scenario success is mistaken for user-validated benefit.','Separate bench, software, scenario and representative-user evidence.')],[1.45,2.52,2.78])
h('Remaining schedule')
table(['Planned milestone','Work to close','Actual completion'],[
('25–26 September','Integration, priority fault checks and evidence capture.',''),
('27 September','Validation and content freeze from the original plan.',''),
('28–29 September','Report evidence audit, poster artwork and submission review.',''),
('30 September','Project report and display poster artwork due.',''),
('7 October','25-minute pre-event prototype showcase and demonstration.',''),
('15 October','Main-event slideshow and video submission due.',''),
('24 October','Main event: 5-minute talk, up to 5-minute video and 5-minute Q&A.','')],[1.48,3.62,1.65])
note('25–29 September entries are internal planning. Official deadlines and event timing follow the 18 August update. Event dates/venues remain tentative. The 24 October presentation applies if selected for the main event. [1, 11, pp. 14–18]')

page('10 Proposed budget and team records')
p('The following sections are reserved for the group to complete after decisions and trials. Empty fields should remain empty until the associated record or result is available.')
h('Final team responsibilities')
table(['Role from the action plan','Member','Final contribution'],[
('Systems and software lead','',''),('ML and computer vision lead','',''),('Stop unit and sensing lead','',''),('Bus electronics and control lead','',''),('Mechanical and validation lead','','')],[2.55,1.6,2.6])
h('Proposed prototype budget')
table(['Item or category','Quantity','Unit cost SGD','Proposed total SGD'],[
('Controllers and communications','','',''),('Camera and VL53L0X sensor','','',''),('Ramp actuator and driver','','',''),('Switches and feedback devices','','',''),('Power and wiring','','',''),('Chassis and boarding fixture','','',''),('Contingency','','',''),('Total proposed budget','','','')],[2.7,.8,1.62,1.63])
note('Complete quantities, quotations and totals before submission. Seed funding amount, in-kind items and actual expenditure should be recorded separately.')
blank('User and stakeholder feedback',0)
h('Conclusion')
p('SG GoAssist has established the software structure, mechanical concept and initial ranging setup needed to progress toward an integrated assistance demonstrator. The remaining work is to connect those elements and measure their behavior under nominal, obstructed and failed conditions. Claims about improved independence or passenger satisfaction will depend on the unfinished comparison and feedback work.')
blank('Final conclusion after validation',0)
blank('Final prototype photographs and demonstration link',1)

page('11 References')
p('Project documents and implementation records supporting this draft are listed below. Bracketed references identify the relevant sources. Repository paths are relative to the GoAssist project root.')
refs=[
('1','SG GoAssist Four-Week Prototype Action Plan','SG_GoAssist_Prototype_Action_Plan.pdf. Project window 31 August–30 September 2026. Primary structure, ML workstreams, D1–D9, acceptance targets and schedule.'),
('2','Autonomous Passenger Assistance System Design Document','APAS_Design_Document.docx. Project objective, user needs, proposed physical request and ramp workflow, requirements and fault concepts.'),
('3','Grand Challenge integrated prototype','docs/grand-challenge-integrated-prototype.md. Assistance loop, intent policy, interlocks, persistence, audit and evidence workflow.'),
('4','Feature readiness and project overview','docs/feature-readiness-matrix.md and README.md. Software status, fixtures and physical-device release gates.'),
('5','Requirements and integrated hardware','docs/task-requirements-traceability.md and hardware/INTEGRATED_PROTOTYPE.md. Requirement mapping, mock-bus configuration, safety checks and physical acceptance work.'),
('6','Robobus v2 validation and design exports','cad/VALIDATION_v2.md, cad/README_v2.md, cad/v2_final_report.json and cad/v2_clearance_report.json. Figure 1: cad/Punggol_WeRide_Robobus_v2_deployed.png.'),
('7','ESP32-S3 sensor bring-up and serial evidence','hardware/esp32-s3-tof-test/README.md and artifacts/apas-tof-build/serial-powercycle-test.txt. Verification record dated 23 September 2026.'),
('8','Edge observer and ML 1 package','hardware/edge-observer/README.md, hardware/edge-observer/ml1/README.md, LABEL_GUIDE.md, train_ml1.py and evaluate_ml1.py. Local inference, dataset guidance and reproducible evaluation.'),
('9','Original expression of interest','Singapore BusTech Grand Challenge 2026 - EOI.docx. Initial RFID, camera/LiDAR and pressure/magnetic sensing concept.'),
('10','Circuit and physical setup drawings','artifacts/apas-circuit/APAS_ESP32S3_RYS1230_BC337_Circuit.pdf and output/pdf/APAS_Three_Laser_Physical_Setup.pdf. Proposed wiring and projection setup, separate from the sensor test.'),
('11','Singapore BusTech Grand Challenge 2026 IHL Student Category Updates','Singapore Bus Academy. SGBTGC 2026 - IHL Updates 18 Aug 2026.pdf. 18 August 2026. Pages 8–13: task and success criteria. Pages 14–17: events, rubric and deadlines. Page 18: report content, format and submission requirements.'),
('12','VL53L0X sensor specifications','STMicroelectronics. VL53L0X Time-of-Flight ranging sensor. Datasheet DS11555, revision 6, June 2024. https://www.st.com/resource/datasheet/vl53l0x.pdf. Product overview: https://www.st.com/en/imaging-and-photonics-solutions/vl53l0x. Accessed 25 September 2026.')]
for num,title,detail in refs:
    a=p(); a.add_run(num+'  '+title).bold=True
    a.paragraph_format.keep_with_next=True
    b=p(detail); b.paragraph_format.space_after=Pt(12)
    for r in b.runs: r.font.size=Pt(10)

path=OUT/'NUS - SmaRHt Buses - Project Report Draft.docx'
doc.save(path)
print(path)
